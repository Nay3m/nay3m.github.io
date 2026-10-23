from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FONT = "Arial"
INK = RGBColor(31, 41, 51)
FOREST = RGBColor(24, 53, 46)
MOSS = RGBColor(73, 107, 84)
MUTED = RGBColor(86, 97, 105)


def set_font(run, size, color=INK, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style(style, size, color=INK, bold=False, italic=False):
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    style.font.italic = italic
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0


def format_paragraph(paragraph, size, color=INK, bold=None, italic=None, before=0, after=0, left=0, first=0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.0
    fmt.left_indent = Inches(left)
    fmt.first_line_indent = Inches(first)
    fmt.keep_together = True
    for run in paragraph.runs:
        set_font(run, size, color=color, bold=bold, italic=italic)


path = Path("Jannatul-Nayem-CV.docx")
document = Document(path)
section = document.sections[0]
section.top_margin = Inches(0.48)
section.bottom_margin = Inches(0.48)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)

set_style(document.styles["Normal"], 9.2)
for style_name in ("Body Text", "Entry"):
    if style_name in document.styles:
        set_style(document.styles[style_name], 9.2)
if "CV Bullet" in document.styles:
    set_style(document.styles["CV Bullet"], 9.0)
if "Section Heading" in document.styles:
    set_style(document.styles["Section Heading"], 10.5, FOREST, bold=True)
if "Title" in document.styles:
    set_style(document.styles["Title"], 24, FOREST, bold=True)
if "Subtitle" in document.styles:
    set_style(document.styles["Subtitle"], 10.5, MOSS, italic=True)

for paragraph in list(document.paragraphs):
    if not paragraph.text.strip():
        paragraph._element.getparent().remove(paragraph._element)

for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    style_name = paragraph.style.name
    if index == 0:
        format_paragraph(paragraph, 24, color=FOREST, bold=True, before=0, after=1)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif index == 1:
        format_paragraph(paragraph, 10.5, color=MOSS, italic=True, before=0, after=3)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif index in (2, 3):
        format_paragraph(paragraph, 8.7, color=MUTED, before=0, after=1)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif style_name == "Section Heading":
        if text.lower().startswith("field work experience"):
            paragraph.style = document.styles["Body Text"]
            format_paragraph(paragraph, 9.2, color=INK, before=1, after=1, left=0.12)
        else:
            paragraph.text = text.upper()
            paragraph.style = document.styles["Section Heading"]
            format_paragraph(paragraph, 10.5, color=FOREST, bold=True, before=7, after=2)
            paragraph.keep_with_next = True
    elif style_name == "CV Bullet":
        paragraph.style = document.styles["CV Bullet"]
        format_paragraph(paragraph, 9.0, color=INK, before=0, after=1, left=0.18, first=-0.12)
    elif style_name == "Entry":
        paragraph.style = document.styles["Entry"]
        format_paragraph(paragraph, 9.2, color=INK, before=1, after=0.5, left=0.04)
        paragraph.keep_with_next = "\n" in text or "|" in text
    else:
        format_paragraph(paragraph, 9.2, color=INK, before=0, after=1)

for paragraph in document.sections[0].footer.paragraphs:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(paragraph, 8, color=MUTED, before=0, after=0)

document.core_properties.author = "Jannatul Nayem"
document.core_properties.title = "Jannatul Nayem Curriculum Vitae"
document.save(path)
print(f"formatted {path}")
